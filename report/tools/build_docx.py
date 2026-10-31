# -*- coding: utf-8 -*-
"""report.md -> report.docx：转换前变换 + pandoc + 后处理（精简版）"""
import sys, re, subprocess, os, zipfile
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook

REPORT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_MD = os.path.join(REPORT_DIR, 'report.md')
TMP_MD = os.path.join(REPORT_DIR, '.tmp-convert.md')
REF = os.path.join(REPORT_DIR, 'ref-template.docx')
OUT = os.path.join(REPORT_DIR, 'report.docx')

text = open(SRC_MD, encoding='utf-8').read()
# 转换前变换：去一级标题(改用 metadata Title)、去手动目录(docx 用原生 TOC)、
# 图片 alt 置空限宽(避免自动题注与"图 1"重复)、标题去反引号
text = re.sub(r'^# 实验报告\s*\n', '', text, count=1)
text = re.sub(r'^\s*- \[.*\n', '', text, flags=re.M)
# 学号/姓名行（标题后的连续若干行）从正文剥离，后处理时以独立 Author 段插回标题之后（学号在前）
author_lines = []
while True:
    m = re.search(r'^(?:学号|姓名)：.*$', text, re.M)
    if m is None:
        break
    author_lines.append(m.group(0).strip())
    text = text[:m.start()] + text[m.end():]
author_lines.sort(key=lambda l: 0 if l.startswith('学号') else 1)
text = text.replace('![数据结构设计图](数据结构设计图.png)',
                    '![](数据结构设计图.png){width=100%}')
text = text.replace('#### `main.cpp` 说明', '#### main.cpp 说明')
open(TMP_MD, 'w', encoding='utf-8').write(text)

r = subprocess.run(['pandoc', '.tmp-convert.md', '-f', 'markdown', '-t', 'docx',
                    '-o', 'report.docx', '--reference-doc', REF,
                    '--metadata', 'title=实验报告', '--metadata', 'toc-title=目录',
                    '--toc', '--toc-depth=3'],
                   cwd=REPORT_DIR, capture_output=True, text=True, encoding='utf-8', errors='replace')
assert r.returncode == 0, r.stderr

doc = Document(OUT)

# 清除 pandoc 在数学对象后写入的零宽空格（部分字体下渲染为可见方框）
zwsp = 0
for t_el in doc.element.iter(qn('w:t')):
    if t_el.text and '\u200b' in t_el.text:
        t_el.text = t_el.text.replace('\u200b', '')
        zwsp += 1
print('ZWSP cleaned:', zwsp)

# 标题之后、目录之前插入两行作者段（学号/姓名各占一行）
title_p = next(p for p in doc.paragraphs if p.style.name == 'Title')
for line in reversed(author_lines):
    ap = doc.add_paragraph(line, style='Author')
    title_p._p.addnext(ap._p)

def set_ind(pPr, left=None, hanging=None):
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if left is not None:
        ind.set(qn('w:left'), str(left))
    if hanging is not None:
        ind.set(qn('w:hanging'), str(hanging))
    # 显式写零：阻断样式链 firstLineChars=200 穿透（删除属性会让样式值漏进来）
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:firstLineChars'), '0')

def shade(cell, fill='F2F2F2'):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el

# ---- 1) 嵌入 result.xlsx 三块为原生表格（替换 xlsx 超链接段落） ----
wb = load_workbook(os.path.join(REPORT_DIR, 'result.xlsx'), data_only=True)
ws = wb.active

def cell_disp(c):
    v = c.value
    if v is None:
        return ''
    if isinstance(v, float):
        if '%' in (c.number_format or ''):
            return f'{v * 100:.3f}%'
        return f'{v:.6f}'.rstrip('0').rstrip('.')
    return str(v)

FS = Pt(9.5)
FRACTION_RE = re.compile(r'^(\d+)/(\d+)\n\((.+)\)$')
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def add_fraction(p, num, den):
    p._p.append(parse_xml(
        f'<m:oMath xmlns:m="{M_NS}" xmlns:w="{W_NS}"><m:f>'
        f'<m:fPr><m:ctrlPr><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
        f'<w:sz w:val="19"/></w:rPr></m:ctrlPr></m:fPr>'
        f'<m:num><m:r><w:rPr><w:sz w:val="19"/></w:rPr><w:t>{num}</w:t></m:r></m:num>'
        f'<m:den><m:r><w:rPr><w:sz w:val="19"/></w:rPr><w:t>{den}</w:t></m:r></m:den>'
        f'</m:f></m:oMath>'))

anchor = next(p for p in doc.paragraphs
              if p.style.name == 'Heading 3' and p.text.strip() == '实验结果')
link_p = next(p for p in doc.paragraphs if 'result.xlsx' in p.text)
new_tables = []
for base in (2, 26, 50):
    tbl = doc.add_table(rows=22, cols=8)
    tbl.style = doc.styles['Table']
    anchors, inner = set(), set()
    for rng in [g for g in ws.merged_cells.ranges
                if g.min_row >= base and g.max_row <= base + 21]:
        for r_ in range(rng.min_row, rng.max_row + 1):
            for c_ in range(rng.min_col, rng.max_col + 1):
                inner.add((r_ - base, c_ - 1))
        anchors.add((rng.min_row - base, rng.min_col - 1))
        tbl.cell(rng.min_row - base, rng.min_col - 1).merge(
            tbl.cell(rng.max_row - base, rng.max_col - 1))
    for ri in range(22):
        for ci in range(8):
            if (ri, ci) in inner and (ri, ci) not in anchors:
                continue
            txt = cell_disp(ws.cell(row=base + ri, column=ci + 1))
            if not txt:
                continue
            p = tbl.cell(ri, ci).paragraphs[0]
            fm = FRACTION_RE.match(txt)
            if fm:
                add_fraction(p, fm.group(1), fm.group(2))
                p.add_run().add_break()
                p.add_run(f'({fm.group(3)})').font.size = FS
                continue
            run = p.add_run(txt.split('\n')[0])
            run.font.size = FS
            if ri <= 1:
                run.font.bold = True
            for extra in txt.split('\n')[1:]:
                run.add_break()
                p.add_run(extra).font.size = FS
    for ri in range(22):
        for ci in range(8):
            cell = tbl.cell(ri, ci)
            if ri <= 1 or ci == 0:  # 顶部两行与左侧指标列灰底，数据区透明
                shade(cell)
            for extra_p in cell.paragraphs[1:]:
                if not extra_p.text.strip():
                    extra_p._p.getparent().remove(extra_p._p)
    new_tables.append(tbl)
link_p._p.getparent().remove(link_p._p)
last = anchor._p
for tbl in new_tables:  # 相邻 tbl 会被 Word 视为一张表，用小空段隔开
    spacer = doc.add_paragraph()
    spacer.add_run().font.size = Pt(6)
    set_ind(spacer._p.get_or_add_pPr())
    last.addnext(tbl._tbl)
    tbl._tbl.addnext(spacer._p)
    last = spacer._p
print('embedded result tables:', len(new_tables))

# ---- 2) 表格统一处理：列宽按内容自适应（CJK=2/ASCII=1 单位，取各列最大内容宽），整表与单元格居中 ----
PAGE_TEXT_TWIPS = 8465
MIN_COL_TWIPS = 850    # 最窄列兜底（含边距后仍可容纳 O(n) 类短值）
UNIT_TWIPS = 105       # 每单位约合 0.5 个汉字宽
MAX_COL_UNITS = 46     # 单列权重上限，防超长说明独占全表

def cell_units(cell):
    # 公式内容在 m:t 中，与 w:t 一并计量
    text = ''.join(re.findall(r'<(?:w|m):t[^>]*>([^<]*)</(?:w|m):t>', cell._tc.xml))
    units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
    return max(min(units, MAX_COL_UNITS), 4)

for tbl in doc.tables:
    look = get_or_add(tbl._tbl.find(qn('w:tblPr')), 'w:tblLook')
    look.set(qn('w:firstRow'), '1')
    for cell in tbl.rows[0].cells:
        # 表头行跨页自动重复
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        if trPr.find(qn('w:tblHeader')) is None:
            trPr.append(OxmlElement('w:tblHeader'))
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn('w:tcBorders'))
        if borders is None:
            borders = OxmlElement('w:tcBorders')
            tcPr.append(borders)
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single')
        bot.set(qn('w:sz'), '6')
        bot.set(qn('w:space'), '0')
        borders.append(bot)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
    # 列宽自适应：各列取最大内容单位宽，按比例分配页心
    ncols = len(tbl._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')))
    natural = [0.0] * ncols
    for row in tbl.rows:
        seen_tcs = set()
        for ci, cell in enumerate(row.cells):
            if id(cell._tc) in seen_tcs:   # 合并单元格跨列时只计一次
                continue
            seen_tcs.add(id(cell._tc))
            u = cell_units(cell)
            if u > natural[ci]:
                natural[ci] = u
    twips = [max(u * UNIT_TWIPS, MIN_COL_TWIPS) for u in natural]
    scale = PAGE_TEXT_TWIPS / sum(twips)
    widths = [max(int(round(w * scale)), MIN_COL_TWIPS) for w in twips]
    widths[-1] += PAGE_TEXT_TWIPS - sum(widths)
    cols = tbl._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol'))
    for c, w in zip(cols, widths):
        c.set(qn('w:w'), str(w))
    n_cols = len(cols)
    for row in tbl.rows:
        # 行不跨页拆分（尾部内容被切到下一页而表头不跟随，是"内容缺失"观感的来源）
        # cantSplit 必须位于 tblHeader 之前（trPr 子元素顺序），故插到最前
        trPr = row._tr.get_or_add_trPr()
        trPr.insert(0, OxmlElement('w:cantSplit'))
        for ci, (cell, w) in enumerate(zip(row.cells, widths)):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = get_or_add(tcPr, 'w:tcW')   # 必须复用 pandoc 已写入的 tcW，追加重复会失效
            tcW.set(qn('w:w'), str(w))
            tcW.set(qn('w:type'), 'dxa')
            get_or_add(tcPr, 'w:vAlign').set(qn('w:val'), 'center')
            # 短值列居中；说明列（3/5 列表的最后一列）长文本左对齐
            align = WD_ALIGN_PARAGRAPH.CENTER
            if ci == n_cols - 1 and n_cols in (3, 5):
                align = WD_ALIGN_PARAGRAPH.LEFT
            for p in cell.paragraphs:
                p.alignment = align
                # 表格内容禁用首行缩进（Compact 继承 BodyText 的 2 字符缩进，把单元格啃掉一大截）
                set_ind(p._p.get_or_add_pPr())
    tblPr = tbl._tbl.find(qn('w:tblPr'))
    # OOXML 要求 tblPr 子元素顺序 tblW → jc → tblLayout（早于 tblLook），乱序会触发 Word 文件修复丢内容
    jc = get_or_add(tblPr, 'w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.remove(jc)
    tblW = get_or_add(tblPr, 'w:tblW')
    tblW.set(qn('w:w'), str(PAGE_TEXT_TWIPS))
    tblW.set(qn('w:type'), 'dxa')
    if tblPr.find(qn('w:tblLayout')) is None:
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        tblW.addnext(layout)
    else:
        layout = tblPr.find(qn('w:tblLayout'))
    # 收紧默认单元格左右边距（108→45 twips），表格两侧观感空间大主要是这个默认值
    if tblPr.find(qn('w:tblCellMar')) is None:
        cellMar = OxmlElement('w:tblCellMar')
        for side, wv in (('w:top', '20'), ('w:left', '45'), ('w:bottom', '20'), ('w:right', '45')):
            el = OxmlElement(side)
            el.set(qn('w:w'), wv)
            el.set(qn('w:type'), 'dxa')
            cellMar.append(el)
        layout.addnext(cellMar)
    tblW.addnext(jc)
print('tables processed:', len(doc.tables))

# ---- 3) 图片与题注居中 ----
center_next = False
for p in doc.paragraphs:
    is_img = bool(p._p.findall('.//' + qn('w:drawing')))
    if is_img or (center_next and p.text.strip()):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center_next = is_img

# ---- 4) 列表缩进约定：序号起点每级 2 个汉字，续行平行对齐；suff=nothing 绕开制表位 ----
CHAR, MARKER = 220, 240
IND = {0: CHAR * 2 + MARKER, 1: CHAR * 4 + MARKER}
last_ilvl = None
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        last_ilvl = None
        continue
    pPr = p._p.find(qn('w:pPr'))
    numPr = pPr.find(qn('w:numPr')) if pPr is not None else None
    if numPr is not None:
        ilvl_el = numPr.find(qn('w:ilvl'))
        ilvl = int(ilvl_el.get(qn('w:val'))) if ilvl_el is not None else 0
        set_ind(pPr, left=IND[ilvl], hanging=MARKER)
        last_ilvl = ilvl
    elif p.style.name in ('Body Text', 'First Paragraph') and last_ilvl is not None:
        set_ind(pPr, left=IND[last_ilvl])
print('list & continuation paragraphs normalized')

doc.save(OUT)

# ---- 5) numbering.xml 同步缩进 + suff=nothing；settings.xml 打开时自动更新域 ----
z = zipfile.ZipFile(OUT)
items = {i.filename: z.read(i.filename) for i in z.infolist()}
z.close()
numbering = items['word/numbering.xml'].decode('utf-8')
for m in list(re.finditer(r'<w:lvl w:ilvl="([01])"[^>]*>.*?</w:lvl>', numbering, re.S)):
    left = IND[int(m.group(1))]
    block = re.sub(r'<w:ind [^/]*/>', f'<w:ind w:left="{left}" w:hanging="{MARKER}"/>', m.group(0))
    numbering = numbering.replace(m.group(0), block)
numbering = numbering.replace('<w:numFmt w:val="decimal" />',
                              '<w:numFmt w:val="decimal" /><w:suff w:val="nothing"/>')
numbering = numbering.replace('<w:numFmt w:val="decimal"/>',
                              '<w:numFmt w:val="decimal"/><w:suff w:val="nothing"/>')
items['word/numbering.xml'] = numbering.encode('utf-8')
settings = items['word/settings.xml'].decode('utf-8')
if 'updateFields' not in settings:
    items['word/settings.xml'] = re.sub(r'(</w:settings>)',
                                        '<w:updateFields w:val="true"/></w:settings>',
                                        settings, count=1).encode('utf-8')
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for name, data in items.items():
        zout.writestr(name, data)

os.remove(TMP_MD)
print('DONE')
