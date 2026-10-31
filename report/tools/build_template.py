# -*- coding: utf-8 -*-
"""从 pandoc 默认模板生成 ref-template.docx（注意：会覆盖手动修改，慎用）"""
import sys, os, subprocess
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.dirname(TOOLS_DIR)
SRC = os.path.join(TOOLS_DIR, '.ref-default.docx')
DST = os.path.join(REPORT_DIR, 'ref-template.docx')
open(SRC, 'wb').write(subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'],
                                     capture_output=True).stdout)

doc = Document(SRC)
styles = doc.styles

def set_fonts(style, ascii_font, ea_font='等线'):
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rFonts.get(qn(attr)) is not None:
            del rFonts.attrib[qn(attr)]
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), ea_font)

def black(style):
    style.font.color.rgb = RGBColor(0, 0, 0)

def pf(style, line=None, before=None, after=None, align=None, keep=None):
    p = style.paragraph_format
    if line is not None:
        p.line_spacing = line
    if before is not None:
        p.space_before = Pt(before)
    if after is not None:
        p.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if keep is not None:
        p.keep_with_next = keep

def first_line_chars(style):
    pPr = style.element.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLine'), '440')
    ind.set(qn('w:firstLineChars'), '200')
    pPr.append(ind)

# 正文族：等线 11pt 纯黑，正文首行缩进 2 字符 + 1.5 倍行距
for name in ('Normal', 'Body Text', 'First Paragraph'):
    s = styles[name]
    set_fonts(s, '等线')
    s.font.size = Pt(11)
    black(s)
pf(styles['Body Text'], line=1.5, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
first_line_chars(styles['Body Text'])
pf(styles['First Paragraph'], line=1.5, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
first_line_chars(styles['First Paragraph'])
pf(styles['Compact'], line=1.3, after=0)

# 代码样式（pandoc 默认模板缺失，需新建）
code = styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
code.style_id = 'SourceCode'
code.base_style = styles['Normal']
set_fonts(code, 'Consolas')
code.font.size = Pt(9.5)
black(code)
pf(code, line=1.0, before=6, after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
pPr = code.element.get_or_add_pPr()
shd = OxmlElement('w:shd')
shd.set(qn('w:val'), 'clear')
shd.set(qn('w:color'), 'auto')
shd.set(qn('w:fill'), 'F2F2F2')
pPr.append(shd)

verb = styles['Verbatim Char']
set_fonts(verb, 'Consolas')
# 补 cs 槽位：等线没有真正的半角字形，ASCII 落入 eastAsia 槽会被画成近全角，表格内代码名因此超宽
rPr_verb = verb.element.get_or_add_rPr()
rFonts_verb = rPr_verb.find(qn('w:rFonts'))
rFonts_verb.set(qn('w:cs'), 'Consolas')
rFonts_verb.set(qn('w:hint'), 'default')
verb.font.size = Pt(10)
black(verb)

# 标题族：等线加粗纯黑、显式关斜体（pandoc 默认 H4 自带斜体）
title = styles['Title']
set_fonts(title, '等线')
title.font.size = Pt(22)
title.font.bold = True
black(title)
pf(title, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)

for name, (size, before, after) in {'Heading 1': (16, 24, 8), 'Heading 2': (15, 18, 8),
                                    'Heading 3': (13, 14, 6), 'Heading 4': (12, 10, 4),
                                    'Heading 5': (11, 8, 4)}.items():
    h = styles[name]
    set_fonts(h, '等线')
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.italic = False
    black(h)
    pf(h, line=1.2, before=before, after=after, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)

toc_h = styles['TOC Heading']
set_fonts(toc_h, '等线')
toc_h.font.size = Pt(15)
toc_h.font.bold = True
black(toc_h)

# 作者区（学号姓名）：等线，避免全角空格在主题字体下变方框
author_style = styles['Author']
set_fonts(author_style, '等线')
author_style.font.size = Pt(12)
black(author_style)
pf(author_style, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)

# 三线表：上下粗框 + 行间细横线，无竖线；表头条件格式加粗+底线
tbl = styles['Table']
tblPr = tbl.element.find(qn('w:tblPr'))
for old in tblPr.findall(qn('w:tblBorders')):
    tblPr.remove(old)
borders = OxmlElement('w:tblBorders')
for tag, val, sz in (('w:top', 'single', '12'), ('w:bottom', 'single', '12'),
                     ('w:insideH', 'single', '4'), ('w:left', 'none', '0'),
                     ('w:right', 'none', '0'), ('w:insideV', 'none', '0')):
    el = OxmlElement(tag)
    el.set(qn('w:val'), val)
    el.set(qn('w:sz'), sz)
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), 'auto')
    borders.append(el)
tblPr.append(borders)
firstRow = tbl.element.find(qn('w:tblStylePr'))
rPr = firstRow.find(qn('w:rPr'))
if rPr is None:
    rPr = OxmlElement('w:rPr')
    firstRow.insert(0, rPr)
rPr.append(OxmlElement('w:b'))
tcBorders = firstRow.find(qn('w:tcPr')).find(qn('w:tcBorders'))
bot = tcBorders.find(qn('w:bottom'))
bot.set(qn('w:sz'), '6')
bot.set(qn('w:space'), '0')

# 超链接固定蓝（不随 Word 主题变）
hl = styles['Hyperlink']
color = hl.element.get_or_add_rPr().find(qn('w:color'))
if color is None:
    color = OxmlElement('w:color')
    hl.element.get_or_add_rPr().append(color)
for attr in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
    if color.get(qn(attr)) is not None:
        del color.attrib[qn(attr)]
color.set(qn('w:val'), '0563C1')

# A4 + 页边距 + 页脚居中页码
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)
sec.top_margin = sec.bottom_margin = Cm(2.54)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
run._r.append(begin); run._r.append(instr); run._r.append(end)
run.font.size = Pt(9)

doc.save(DST)
os.remove(SRC)
print('template saved:', DST)
